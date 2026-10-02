from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt
from PIL import Image

ROOT = Path(__file__).resolve().parent
OUTPUT = ROOT / "WeatherGrid_Final_Presentation.pptx"

NAVY = RGBColor(19, 51, 68)
TEAL = RGBColor(4, 127, 140)
BLUE = RGBColor(43, 108, 176)
GREEN = RGBColor(30, 130, 94)
PURPLE = RGBColor(111, 78, 176)
AMBER = RGBColor(185, 108, 15)
RED = RGBColor(190, 62, 72)
INK = RGBColor(31, 50, 61)
MUTED = RGBColor(88, 108, 118)
LINE = RGBColor(216, 228, 233)
PALE = RGBColor(246, 249, 250)
WHITE = RGBColor(255, 255, 255)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank = prs.slide_layouts[6]


def background(slide, color=PALE):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color


def text(slide, x, y, w, h, value, size=18, color=INK, bold=False,
         align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.vertical_anchor = valign
    p = frame.paragraphs[0]
    p.text = value
    p.alignment = align
    p.font.name = "Aptos"
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = color
    return shape


def header(slide, number, title, subtitle=""):
    text(slide, 0.6, 0.3, 0.55, 0.35, f"{number:02}", 12, TEAL, True)
    text(slide, 1.15, 0.2, 11.4, 0.5, title, 26, NAVY, True)
    if subtitle:
        text(slide, 1.15, 0.72, 11.2, 0.35, subtitle, 12, MUTED)
    line = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0.6), Inches(1.14), Inches(12.1), Inches(0.025))
    line.fill.solid(); line.fill.fore_color.rgb = LINE; line.line.fill.background()


def footer(slide, number):
    text(slide, 0.62, 7.14, 4.0, 0.2, "ICT304 · WeatherGrid", 8, MUTED)
    text(slide, 12.0, 7.14, 0.65, 0.2, str(number), 8, MUTED, align=PP_ALIGN.RIGHT)


def box(slide, x, y, w, h, title, body, accent=TEAL, title_size=17, body_size=14):
    card = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    card.fill.solid(); card.fill.fore_color.rgb = WHITE
    card.line.color.rgb = LINE
    stripe = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(x), Inches(y), Inches(0.07), Inches(h))
    stripe.fill.solid(); stripe.fill.fore_color.rgb = accent; stripe.line.fill.background()
    text(slide, x + 0.25, y + 0.2, w - 0.45, 0.4, title, title_size, NAVY, True)
    text(slide, x + 0.25, y + 0.73, w - 0.45, h - 0.9, body, body_size, MUTED)


def bullets(slide, x, y, w, h, values, size=18, color=INK):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.clear(); frame.word_wrap = True
    for index, value in enumerate(values):
        p = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        p.text = "•  " + value
        p.font.name = "Aptos"
        p.font.size = Pt(size)
        p.font.color.rgb = color
        p.space_after = Pt(14)
    return shape


def image_fit(slide, path, x, y, w, h):
    with Image.open(path) as source:
        iw, ih = source.size
    scale = min(w / iw, h / ih)
    pw, ph = iw * scale, ih * scale
    return slide.shapes.add_picture(str(path), Inches(x + (w - pw) / 2), Inches(y + (h - ph) / 2), Inches(pw), Inches(ph))


def pill(slide, x, y, w, value, color=TEAL):
    shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(0.43))
    shape.fill.solid(); shape.fill.fore_color.rgb = color; shape.line.fill.background()
    text(slide, x, y + 0.02, w, 0.34, value, 10, WHITE, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)


def notes(slide, value):
    slide.notes_slide.notes_text_frame.text = value


# 1 Title
slide = prs.slides.add_slide(blank); background(slide, NAVY)
text(slide, 0.8, 0.65, 3.5, 0.35, "ICT304 DISTRIBUTED COMPUTING", 12, RGBColor(119, 220, 211), True)
text(slide, 0.8, 1.45, 9.6, 0.85, "WeatherGrid", 42, WHITE, True)
text(slide, 0.8, 2.35, 10.6, 0.55, "Distributed IoT Weather Monitoring System", 25, RGBColor(211, 232, 238), True)
text(slide, 0.8, 3.3, 8.6, 0.7, "Local weather data delivered reliably from simulated devices to a live web dashboard.", 19, WHITE)
pill(slide, 0.8, 4.38, 1.65, "3 stations", BLUE)
pill(slide, 2.62, 4.38, 1.75, "MQTT messaging", PURPLE)
pill(slide, 4.54, 4.38, 1.75, "Edge gateway", TEAL)
pill(slide, 6.46, 4.38, 1.75, "Live dashboard", GREEN)
text(slide, 0.8, 5.55, 8.5, 0.75, "Group [NUMBER]\n[ADD FIVE MEMBER NAMES AND STUDENT IDs]", 14, RGBColor(211, 232, 238))
text(slide, 10.5, 6.7, 1.9, 0.25, "15-minute presentation", 9, RGBColor(166, 197, 207), align=PP_ALIGN.RIGHT)
notes(slide, "Member 1, 15 seconds: Good morning. We are presenting WeatherGrid, a distributed IoT weather monitoring system. It collects local weather data and shows it on a live web dashboard.")

# 2 Problem
slide = prs.slides.add_slide(blank); background(slide); header(slide, 2, "The problem", "Why did we build WeatherGrid?")
box(slide, 0.75, 1.55, 3.75, 4.65, "Local weather changes", "Public forecasts cover large areas. A farm, campus or workplace may need weather information for its exact location.", RED, 18, 16)
box(slide, 4.78, 1.55, 3.75, 4.65, "Devices can lose connection", "Weather devices may disconnect. Messages can be delayed, lost or sent more than once.", PURPLE, 18, 16)
box(slide, 8.56, 1.55, 3.75, 4.65, "Users need clear information", "People need live readings, warnings, history and reports in one simple and secure application.", TEAL, 18, 16)
text(slide, 1.0, 6.52, 11.3, 0.35, "Our goal: reliable local monitoring that continues to work during short failures.", 18, NAVY, True, PP_ALIGN.CENTER)
footer(slide, 2)
notes(slide, "Member 1, 45 seconds: Explain these three problems in simple words. Finish by saying that reliability is important because devices and networks are not always available.")

# 3 Objectives
slide = prs.slides.add_slide(blank); background(slide); header(slide, 3, "Project objectives", "What should the system do?")
bullets(slide, 0.9, 1.55, 5.7, 4.8, [
    "Collect six weather measurements",
    "Support several weather stations",
    "Show live and historical data",
    "Detect warning and critical conditions",
], 19)
bullets(slide, 6.8, 1.55, 5.6, 4.8, [
    "Avoid duplicate database records",
    "Keep messages during short outages",
    "Protect admin and viewer access",
    "Provide reports and a test laboratory",
], 19)
pill(slide, 3.95, 6.15, 5.4, "Working prototype + testing + live demonstration", NAVY)
footer(slide, 3)
notes(slide, "Member 1, 30 seconds: Present the objectives as outcomes. Point out that the system must remain reliable when messages are repeated or delayed.")

# 4 Architecture
slide = prs.slides.add_slide(blank); background(slide); header(slide, 4, "System design", "Seven parts work together as one distributed system")
image_fit(slide, ROOT / "weathergrid-high-level-architecture-generated.png", 0.45, 1.25, 12.45, 5.75)
footer(slide, 4)
notes(slide, "Member 2, 90 seconds: Move from left to right. Simulators create data. Mosquitto routes MQTT messages. The gateway validates and forwards them. Next.js and Payload process the data. PostgreSQL stores it. Workers check offline devices and notifications. The browser displays the result.")

# 5 Data flow
slide = prs.slides.add_slide(blank); background(slide); header(slide, 5, "How one reading moves through the system", "A simple start-to-end flow")
image_fit(slide, ROOT / "weathergrid-telemetry-flow.png", 0.45, 1.35, 12.45, 5.2)
pill(slide, 4.2, 6.55, 4.9, "Temporary failure → save in SQLite → retry later", PURPLE)
footer(slide, 5)
notes(slide, "Member 2, 60 seconds: Explain the seven numbered steps. Mention that the gateway saves data in SQLite if the backend is unavailable. The browser receives a live update after storage.")

# 6 Prototype and sensors
slide = prs.slides.add_slide(blank); background(slide); header(slide, 6, "Our working prototype", "Three simulated stations using a realistic IoT message format")
box(slide, 0.7, 1.5, 3.7, 2.0, "Sydney", "WX-SYD-001\nIndependent Python process", BLUE, 20, 15)
box(slide, 4.8, 1.5, 3.7, 2.0, "Melbourne", "WX-MEL-001\nIndependent Python process", PURPLE, 20, 15)
box(slide, 8.9, 1.5, 3.7, 2.0, "Brisbane", "WX-BNE-001\nIndependent Python process", TEAL, 20, 15)
text(slide, 0.85, 4.05, 11.7, 0.45, "Temperature · Humidity · Pressure · Rainfall · Wind speed · Battery", 21, NAVY, True, PP_ALIGN.CENTER)
text(slide, 1.05, 4.95, 11.2, 0.95, "The current weather values are simulated. The timestamps are real. Future ESP32 devices can use BME280/DHT22, a rain gauge, an anemometer and a battery sensor.", 17, MUTED, align=PP_ALIGN.CENTER)
pill(slide, 4.55, 6.17, 4.25, "Same message format for simulated or physical devices", GREEN)
footer(slide, 6)
notes(slide, "Member 3, 45 seconds: Be clear that the sensors are simulated. Explain that each station runs separately and publishes realistic test values. Physical sensors can replace the simulator without changing the backend message contract.")

# 7 Dashboard
slide = prs.slides.add_slide(blank); background(slide); header(slide, 7, "Web application", "Live information for viewers and control tools for administrators")
image_fit(slide, ROOT / "weathergrid-live-monitoring.png", 0.5, 1.3, 8.45, 5.65)
box(slide, 9.15, 1.45, 3.45, 1.35, "Monitor", "Live readings, history and map", BLUE, 16, 12)
box(slide, 9.15, 3.0, 3.45, 1.35, "Respond", "Alerts, device status and notifications", RED, 16, 12)
box(slide, 9.15, 4.55, 3.45, 1.35, "Manage", "Reports, users and simulation commands", TEAL, 16, 12)
pill(slide, 9.15, 6.18, 3.45, "Responsive web application", NAVY)
footer(slide, 7)
notes(slide, "Member 3, 30 seconds: This is a screenshot of the running prototype. Explain the units: Celsius, percent, hectopascals, millimetres, kilometres per hour and battery percent.")

# 8 Reliability/security
slide = prs.slides.add_slide(blank); background(slide); header(slide, 8, "Why the system is reliable and secure", "Important distributed computing decisions")
box(slide, 0.75, 1.5, 5.8, 2.05, "Reliable delivery", "MQTT QoS 1 improves delivery. SQLite keeps messages when the backend is temporarily unavailable.", PURPLE, 19, 15)
box(slide, 6.82, 1.5, 5.8, 2.05, "No duplicate readings", "Device ID and sequence number identify the same message, so retries are stored only once.", AMBER, 19, 15)
box(slide, 0.75, 3.9, 5.8, 2.05, "Validation", "Pydantic and Zod check fields, ranges, IDs and timestamps before data is accepted.", TEAL, 19, 15)
box(slide, 6.82, 3.9, 5.8, 2.05, "Access control", "Admins manage the system. Viewers can monitor it. The gateway uses a separate machine account.", BLUE, 19, 15)
footer(slide, 8)
notes(slide, "Member 4, 60 seconds: Define duplicate prevention simply: sending the same reading again does not create another database row. Mention that production MQTT would require TLS and per-device credentials.")

# 9 Testing/findings
slide = prs.slides.add_slide(blank); background(slide); header(slide, 9, "Testing and findings", "Evidence that the prototype works")
values = [("Web logic", 8, BLUE), ("Gateway", 7, TEAL), ("Simulator", 27, PURPLE)]
for i, (label, count, color) in enumerate(values):
    y = 1.6 + i * 1.15
    text(slide, 0.85, y + 0.1, 1.65, 0.35, label, 15, NAVY, True)
    track = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(2.55), Inches(y), Inches(5.2), Inches(0.58))
    track.fill.solid(); track.fill.fore_color.rgb = LINE; track.line.fill.background()
    bar = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(2.55), Inches(y), Inches(5.2 * count / 30), Inches(0.58))
    bar.fill.solid(); bar.fill.fore_color.rgb = color; bar.line.fill.background()
    text(slide, 7.95, y + 0.1, 1.2, 0.35, f"{count} passed", 14, color, True)
text(slide, 0.85, 5.25, 7.6, 0.48, "42 automated tests passed", 26, NAVY, True)
pill(slide, 0.85, 6.0, 1.75, "TypeScript passed", GREEN)
pill(slide, 2.8, 6.0, 1.55, "ESLint passed", GREEN)
pill(slide, 4.55, 6.0, 2.45, "Services healthy", GREEN)
box(slide, 9.25, 1.55, 3.15, 4.55, "Main findings", "• Three devices send data together\n\n• Live values update correctly\n\n• Alerts respond to test scenarios\n\n• Duplicate and outage handling works\n\n• Data remains available for history and reports", AMBER, 19, 14)
footer(slide, 9)
notes(slide, "Member 4, 45 seconds: Explain that tests include bad data and failure cases, not only normal readings. State the total of 42 tests and show that all main services were healthy.")

# 10 Teamwork
slide = prs.slides.add_slide(blank); background(slide); header(slide, 10, "How our five-person team worked", "Clear ownership with shared integration and testing")
roles = [
    ("Member 1", "Leadership, requirements, architecture and report", BLUE),
    ("Member 2", "Device simulator and MQTT communication", PURPLE),
    ("Member 3", "Edge gateway, buffering and retry", TEAL),
    ("Member 4", "Backend, database, security and alerts", AMBER),
    ("Member 5", "Frontend, testing and demonstration", GREEN),
]
for i, (member, role, color) in enumerate(roles):
    y = 1.43 + i * 1.02
    pill(slide, 0.8, y, 1.55, member, color)
    text(slide, 2.65, y + 0.03, 9.55, 0.4, role, 16, INK, True)
text(slide, 0.9, 6.65, 11.5, 0.32, "Shared work: integration, testing, documentation, rehearsal and problem solving", 16, NAVY, True, PP_ALIGN.CENTER)
footer(slide, 10)
notes(slide, "Member 5, 30 seconds: Replace Member 1 to Member 5 with real names. Each person should explain their own work. State that integration and testing were shared by the whole team.")

# 11 Demo
slide = prs.slides.add_slide(blank); background(slide, NAVY)
text(slide, 0.75, 0.45, 11.8, 0.5, "LIVE PROTOTYPE DEMONSTRATION", 28, WHITE, True)
text(slide, 0.75, 1.05, 10.4, 0.35, "A clear four-minute demonstration", 14, RGBColor(190, 217, 224))
steps = [
    ("1", "Overview", "Confirm Sydney, Melbourne and Brisbane are online"),
    ("2", "Live monitoring", "Show changing values and explain the units"),
    ("3", "Simulation", "Send Temperature critical to Sydney"),
    ("4", "Alert", "Show the new critical alert and acknowledge it"),
    ("5", "Recovery", "Send Reset to normal and show history/report evidence"),
]
for i, (number, title_value, detail) in enumerate(steps):
    y = 1.63 + i * 0.94
    circle = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, Inches(0.85), Inches(y), Inches(0.52), Inches(0.52))
    circle.fill.solid(); circle.fill.fore_color.rgb = TEAL; circle.line.fill.background()
    text(slide, 0.85, y + 0.07, 0.52, 0.3, number, 12, WHITE, True, PP_ALIGN.CENTER)
    text(slide, 1.62, y - 0.02, 1.85, 0.4, title_value, 18, WHITE, True)
    text(slide, 3.55, y + 0.02, 7.85, 0.42, detail, 14, RGBColor(216, 232, 237))
box(slide, 9.5, 6.0, 2.75, 0.7, "Backup", "Use screenshots if needed", RED, 12, 9)
text(slide, 0.8, 6.72, 8.2, 0.3, "Do not show passwords, API keys or the .env file.", 11, RGBColor(247, 191, 198), True)
notes(slide, "All members, 4 minutes: Member 5 controls the computer. Other members explain the flow while the demo runs. Rehearse the command before class. Keep the architecture image and dashboard screenshot ready as a backup.")

# 12 Conclusion
slide = prs.slides.add_slide(blank); background(slide); header(slide, 12, "Conclusion and future work", "What we achieved and what comes next")
box(slide, 0.75, 1.5, 5.8, 4.75, "What we achieved", "✓ Complete device-to-dashboard flow\n\n✓ Live and historical monitoring\n\n✓ Alerts, reports and admin controls\n\n✓ Buffering and duplicate protection\n\n✓ Tested working prototype", GREEN, 21, 16)
box(slide, 6.82, 1.5, 5.8, 4.75, "Future improvements", "1. Connect calibrated ESP32 sensors\n\n2. Add MQTT over TLS\n\n3. Deploy to cloud infrastructure\n\n4. Test larger numbers of devices\n\n5. Add predictive weather analysis", BLUE, 21, 16)
text(slide, 0.9, 6.55, 11.5, 0.4, "WeatherGrid turns distributed sensor messages into useful and reliable information.", 18, NAVY, True, PP_ALIGN.CENTER)
footer(slide, 12)
notes(slide, "Member 1, 30 seconds: Return to the original problem. State that the objectives were achieved, but be clear that physical sensors and production security are future work.")

# 13 Questions and references
slide = prs.slides.add_slide(blank); background(slide, NAVY)
text(slide, 0.8, 1.15, 11.7, 0.65, "Questions?", 40, WHITE, True, PP_ALIGN.CENTER)
text(slide, 1.4, 2.05, 10.5, 0.45, "Thank you", 23, RGBColor(119, 220, 211), True, PP_ALIGN.CENTER)
text(slide, 1.1, 3.15, 11.1, 0.4, "Key references", 14, WHITE, True, PP_ALIGN.CENTER)
text(slide, 1.4, 3.75, 10.5, 1.7,
     "OASIS Open. (2019). MQTT Version 5.0.\n"
     "Gubbi et al. (2013). Internet of Things: A vision, architectural elements, and future directions.\n"
     "Kodali and Mandal. (2016). IoT based weather station.\n"
     "Satyanarayanan. (2017). The emergence of edge computing.",
     12, RGBColor(199, 220, 226), align=PP_ALIGN.CENTER)
text(slide, 2.0, 6.4, 9.3, 0.35, "WeatherGrid · Distributed IoT Weather Monitoring System", 13, RGBColor(166, 197, 207), True, PP_ALIGN.CENTER)
notes(slide, "Invite questions. Every member should be ready to explain their files, MQTT QoS 1, duplicate prevention, outage handling, user roles and the difference between simulated and physical sensors.")

prs.core_properties.title = "WeatherGrid Final ICT304 Presentation"
prs.core_properties.subject = "15-minute group presentation"
prs.core_properties.author = "WeatherGrid Project Group"
prs.core_properties.keywords = "WeatherGrid, IoT, MQTT, distributed computing"
prs.save(OUTPUT)
print(OUTPUT)
