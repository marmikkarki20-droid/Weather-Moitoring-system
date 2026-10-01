from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parent
OUTPUT = ROOT / "WeatherGrid_ICT304_Assessment_3_Report.docx"
FLOWCHART = ROOT / "weathergrid-telemetry-flow.png"

BLUE = "173B57"
TEAL = "087F8C"
LIGHT = "EAF2F7"
PALE = "F4F7F9"
WHITE = "FFFFFF"
GREY = "556575"


def create_flowchart():
    image = Image.new("RGB", (2200, 850), "#F7FAFC")
    draw = ImageDraw.Draw(image)
    font_path = "/System/Library/Fonts/Supplemental/Arial.ttf"
    bold_path = "/System/Library/Fonts/Supplemental/Arial Bold.ttf"
    title_font = ImageFont.truetype(bold_path, 54)
    heading_font = ImageFont.truetype(bold_path, 31)
    body_font = ImageFont.truetype(font_path, 24)
    small_font = ImageFont.truetype(font_path, 21)
    draw.text((75, 45), "WeatherGrid Telemetry Processing Flow", fill="#173B57", font=title_font)
    draw.text((75, 112), "Validated, fault-tolerant path from weather observation to an authenticated live dashboard", fill="#556575", font=body_font)

    boxes = [
        (80, 225, 410, 430, "1  SENSOR NODE", "Sample six measures\nValidate and timestamp", "#E9F2FF", "#2F6FED"),
        (515, 225, 845, 430, "2  MQTT BROKER", "Publish JSON at QoS 1\nRoute by device topic", "#F2ECFF", "#7253D4"),
        (950, 225, 1280, 430, "3  EDGE GATEWAY", "Verify topic and schema\nTimestamp, deliver or buffer", "#EAF7F5", "#087F8C"),
        (1385, 225, 1715, 430, "4  INGESTION API", "Authenticate service\nValidate and transact", "#E8F8EF", "#16835B"),
        (1760, 560, 2090, 765, "5  POSTGRESQL", "Persist once\nUpdate status and alerts", "#FFF7DE", "#B66A00"),
        (1160, 560, 1490, 765, "6  SSE HUB", "Publish change event\nFan out to clients", "#FFF0F1", "#C33B55"),
        (560, 560, 890, 765, "7  WEB DASHBOARD", "Refresh current data\nDisplay values and units", "#E9F2FF", "#2F6FED"),
    ]
    for x1, y1, x2, y2, heading, detail, fill, outline in boxes:
        draw.rounded_rectangle((x1, y1, x2, y2), radius=24, fill=fill, outline=outline, width=5)
        draw.text((x1 + 24, y1 + 27), heading, fill="#152638", font=heading_font)
        draw.multiline_text((x1 + 24, y1 + 91), detail, fill="#394B5A", font=small_font, spacing=10)

    def arrow(start, end, color="#34657F"):
        draw.line((start, end), fill=color, width=10)
        ex, ey = end
        if abs(end[0] - start[0]) > abs(end[1] - start[1]):
            points = [(ex, ey), (ex - 25 if ex > start[0] else ex + 25, ey - 17), (ex - 25 if ex > start[0] else ex + 25, ey + 17)]
        else:
            points = [(ex, ey), (ex - 17, ey - 25 if ey > start[1] else ey + 25), (ex + 17, ey - 25 if ey > start[1] else ey + 25)]
        draw.polygon(points, fill=color)

    arrow((410, 327), (515, 327))
    arrow((845, 327), (950, 327))
    arrow((1280, 327), (1385, 327))
    draw.line((1715, 327, 1925, 327, 1925, 560), fill="#34657F", width=10)
    draw.polygon([(1925, 560), (1908, 530), (1942, 530)], fill="#34657F")
    arrow((1760, 662), (1490, 662))
    arrow((1160, 662), (890, 662))
    draw.text((955, 795), "Temporary API failure -> durable SQLite queue -> bounded retry", fill="#7253D4", font=small_font)
    image.save(FLOWCHART, quality=95)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    tr_pr.append(repeat)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, end])


def add_toc(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = ' TOC \\o "1-3" \\h \\z \\u '
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "Right-click and select Update Field to display the contents."
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for i, header in enumerate(headers):
        cell = hdr.cells[i]
        cell.text = header
        set_cell_shading(cell, BLUE)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for run in cell.paragraphs[0].runs:
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.size = Pt(9)
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = str(value)
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if len(table.rows) % 2 == 0:
                set_cell_shading(cells[i], PALE)
            for p in cells[i].paragraphs:
                p.paragraph_format.space_after = Pt(2)
                for run in p.runs:
                    run.font.size = Pt(8.5)
        if widths:
            for i, width in enumerate(widths):
                cells[i].width = Cm(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.add_run(text)
    return p


def body(doc, text):
    p = doc.add_paragraph(text)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    return p


def code_block(doc, code, caption_text):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_shading(cell, "F1F4F6")
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    for index, line in enumerate(code.splitlines()):
        if index:
            paragraph.add_run("\n")
        run = paragraph.add_run(line)
        run.font.name = "Menlo"
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor.from_string("243746")
    cap = doc.add_paragraph(caption_text, style="Figure Caption")
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return table


create_flowchart()
doc = Document()
section = doc.sections[0]
section.top_margin = Cm(2.2)
section.bottom_margin = Cm(2.0)
section.left_margin = Cm(2.4)
section.right_margin = Cm(2.4)

styles = doc.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"].font.size = Pt(10.5)
styles["Normal"].paragraph_format.space_after = Pt(7)
styles["Normal"].paragraph_format.line_spacing = 1.15
for name, size, color in [("Title", 28, BLUE), ("Heading 1", 17, BLUE), ("Heading 2", 13, TEAL), ("Heading 3", 11, GREY)]:
    styles[name].font.name = "Aptos Display"
    styles[name].font.size = Pt(size)
    styles[name].font.color.rgb = RGBColor.from_string(color)
    styles[name].font.bold = True
    styles[name].font.italic = False
styles["Heading 1"].paragraph_format.page_break_before = True
styles["Heading 1"].paragraph_format.space_after = Pt(8)
styles["Heading 2"].paragraph_format.keep_with_next = True
styles["Heading 3"].paragraph_format.keep_with_next = True

caption = styles.add_style("Figure Caption", WD_STYLE_TYPE.PARAGRAPH)
caption.font.name = "Aptos"
caption.font.size = Pt(9)
caption.font.italic = True
caption.font.color.rgb = RGBColor.from_string(GREY)
caption.paragraph_format.space_after = Pt(8)

for sec in doc.sections:
    add_page_number(sec.footer.paragraphs[0])

# Cover page
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(70)
r = p.add_run("WEATHERGRID")
r.font.name = "Aptos Display"
r.font.size = Pt(34)
r.font.bold = True
r.font.color.rgb = RGBColor.from_string(BLUE)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Distributed IoT Weather Monitoring System")
r.font.size = Pt(19)
r.font.bold = True
r.font.color.rgb = RGBColor.from_string(TEAL)

doc.add_paragraph()
p = doc.add_paragraph("Assessment 3: Group Project Report and Presentation")
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.runs[0].font.size = Pt(15)
p.runs[0].font.bold = True

cover_rows = [
    ("Unit", "ICT304 Distributed Computing"),
    ("Course", "Bachelor of Information Technology (BIT)"),
    ("Semester", "2026 - S1"),
    ("Group Number", "[INSERT GROUP NUMBER]"),
    ("Group Members", "[INSERT NAME] - [INSERT STUDENT ID]\n[INSERT NAME] - [INSERT STUDENT ID]\n[INSERT NAME] - [INSERT STUDENT ID]"),
    ("Submission Date", "[INSERT SUBMISSION DATE]"),
    ("Approximate word count", "3,900 words, excluding tables, captions, references and appendices"),
]
table = doc.add_table(rows=0, cols=2)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.style = "Table Grid"
for label, value in cover_rows:
    cells = table.add_row().cells
    cells[0].text = label
    cells[1].text = value
    set_cell_shading(cells[0], LIGHT)
    cells[0].paragraphs[0].runs[0].font.bold = True
    for cell in cells:
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for para in cell.paragraphs:
            for run in para.runs:
                run.font.size = Pt(10)

doc.add_paragraph()
p = doc.add_paragraph("Academic declaration")
p.runs[0].font.bold = True
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p = doc.add_paragraph("The submitting group confirms that the final document has been reviewed, that all member and contribution details are accurate, and that all sources and reused ideas are acknowledged in APA 7 style.")
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.runs[0].font.size = Pt(9)
doc.add_page_break()

doc.add_heading("Executive Summary", level=1)
body(doc, "WeatherGrid is a distributed Internet of Things (IoT) prototype that collects environmental telemetry from geographically separated weather stations and presents it through a secure, near-real-time web application. The prototype models stations in Sydney, Melbourne and Brisbane. Each station produces temperature, relative humidity, atmospheric pressure, rainfall, wind speed and battery measurements. Messages are validated at the device and edge layers, published using Message Queuing Telemetry Transport (MQTT), persisted in PostgreSQL and delivered to authenticated users through dashboard Application Programming Interfaces (APIs) and Server-Sent Events (SSE).")
body(doc, "The solution addresses three practical distributed-system concerns: components must remain decoupled, duplicate MQTT deliveries must not create duplicate readings, and short application or network outages must not lose accepted device data. WeatherGrid therefore uses Eclipse Mosquitto as a publish/subscribe broker, Quality of Service (QoS) level 1, unique message and sequence identifiers, a durable SQLite gateway buffer with exponential retry, and a database uniqueness constraint. Payload CMS supplies authentication, role-based access control and administrative data management, while Next.js and React provide the responsive dashboard. Automated testing produced 42 passing tests across the web application, gateway and simulator; static type checking and linting also passed. A live verification on 24 September 2026 showed three active devices and more than 4,500 stored readings. The prototype demonstrates a coherent end-to-end distributed IoT architecture and provides a credible base for deployment with physical ESP32 sensor nodes, Transport Layer Security (TLS), managed infrastructure and calibrated sensors.")

doc.add_heading("Table of Contents", level=1)
add_toc(doc.add_paragraph())

doc.add_heading("List of Figures", level=2)
body(doc, "Figure 1. WeatherGrid high-level distributed system architecture\nFigure 2. Telemetry processing flow from observation to live dashboard\nFigure 3. Authenticated live-monitoring prototype")
doc.add_heading("Abbreviations", level=2)
add_table(doc, ["Abbreviation", "Meaning", "Abbreviation", "Meaning"], [
    ("API", "Application Programming Interface", "CMS", "Content Management System"),
    ("HTTP/S", "Hypertext Transfer Protocol / Secure", "IoT", "Internet of Things"),
    ("MQTT", "Message Queuing Telemetry Transport", "QoS", "Quality of Service"),
    ("SSE", "Server-Sent Events", "SQL", "Structured Query Language"),
    ("TLS", "Transport Layer Security", "UUID", "Universally Unique Identifier"),
], [2.4, 5.2, 2.4, 5.2])

doc.add_heading("1. Introduction and Literature Review", level=1)
doc.add_heading("1.1 Project Overview and Problem Statement", level=2)
body(doc, "Local weather can vary substantially across a city or region, while public forecasts often represent a broad area rather than a specific site. Farms, campuses, community organisations and facility operators may therefore need affordable, location-specific observations. A conventional single-machine application is also a poor fit: sensors are distributed, connectivity can be intermittent, multiple users may monitor data concurrently, and measurements must remain consistent despite retries or component failure. The project problem is to design and implement a low-cost distributed platform that reliably gathers local weather observations, stores them centrally, detects abnormal conditions and presents understandable information to authorised users.")
body(doc, "WeatherGrid solves this problem with independently operating weather nodes, brokered messaging, an edge gateway, a transactional application service and a browser-based dashboard. The current academic prototype uses realistic Python sensor simulators so repeatable normal, alert and failure scenarios can be demonstrated without depending on physical hardware. Its contracts are intentionally compatible with future ESP32-based nodes. The result is a webpage rather than a native mobile application, but its responsive layout supports desktop and mobile browsers.")

doc.add_heading("1.2 Project Objectives", level=2)
bullet(doc, "Collect temperature, humidity, pressure, rainfall, wind speed and battery data from multiple independent weather devices.")
bullet(doc, "Transport telemetry through a lightweight publish/subscribe protocol and separate producers from consumers.")
bullet(doc, "Validate, authenticate and persist readings exactly once at the application level despite at-least-once MQTT delivery.")
bullet(doc, "Provide live monitoring, history, device status, mapping, alerts, notifications, reports and controlled simulation through an authenticated web interface.")
bullet(doc, "Tolerate temporary server or network unavailability through durable edge buffering and retry.")
bullet(doc, "Demonstrate quality through reproducible automated, integration and live-system tests.")

doc.add_heading("1.3 Literature Review", level=2)
body(doc, "Gubbi et al. (2013) describe IoT as the convergence of sensing, communications and distributed computing, with sensor data interpreted through shared platforms. Their cloud-centred architecture supports WeatherGrid's separation into device, messaging, edge, application and data layers. However, a cloud-only path can make field devices dependent on continuous wide-area connectivity. WeatherGrid adds an explicit edge gateway and local durable queue so acquisition and central processing are separated in time.")
body(doc, "Kodali and Mandal (2016) demonstrated that a low-cost IoT weather station can collect environmental parameters and expose them remotely. Kodali and Sahu (2016) similarly used a Wi-Fi-enabled WeMos microcontroller to make local weather information accessible through Internet services. These projects validate inexpensive embedded hardware as a practical acquisition layer. WeatherGrid extends the idea beyond a single device by modelling three independent stations and adding central identity, role-based administration, alert lifecycle management, deduplication, historical reporting and outage recovery.")
body(doc, "Satyanarayanan (2017) argues that placing computation near data sources improves responsiveness and can mask transient cloud outages. That principle appears directly in WeatherGrid: the Python gateway validates topic-to-device identity, adds an edge timestamp and retains temporarily undeliverable messages in SQLite. This reduces coupling between MQTT availability and the web application while preserving data for later delivery.")
body(doc, "MQTT is an OASIS standard client/server publish/subscribe transport designed for constrained machine-to-machine and IoT environments (Banks et al., 2019). Its brokered model lets devices publish without knowing the application endpoint. WeatherGrid uses QoS 1 because delivery is more important than avoiding retransmission. QoS 1 can deliver a message more than once, so the application must be idempotent. WeatherGrid addresses this with persistent device sequence numbers, unique message identifiers and a composite database uniqueness constraint. More recent field work by Hernandez et al. (2024) combined meteorological sensing, MQTT and a central visualisation platform, showing the value of integrating collection, processing and dashboards. WeatherGrid follows that pattern while placing stronger emphasis on distributed failure handling, access control and testable simulation. The literature therefore supports the chosen layered architecture, while the prototype contributes a demonstrable reliability path from device publication to user-facing evidence.")

doc.add_heading("2. Requirements Analysis", level=1)
doc.add_heading("2.1 Stakeholders and Scope", level=2)
body(doc, "The primary stakeholders are viewers who monitor conditions, administrators who configure devices and alert rules, and technical operators who maintain the broker, gateway, application and database. The prototype scope includes simulated sensing, messaging, ingestion, visualisation, alerting and reporting. Physical sensor calibration, public cloud deployment and operational forecasting are deliberately outside the current scope; these are future extensions rather than claims of the prototype.")

doc.add_heading("2.2 Functional Requirements", level=2)
add_table(doc, ["ID", "Requirement", "Justification / acceptance condition"], [
    ("FR1", "Acquire six telemetry measures for three stations.", "Every payload contains temperature, humidity, pressure, rainfall, wind speed and battery plus device/time metadata."),
    ("FR2", "Publish telemetry and status using MQTT.", "Each device uses its own topic; telemetry uses QoS 1 and online/offline status is retained."),
    ("FR3", "Validate at device, gateway and API boundaries.", "Malformed, oversized, out-of-range, mismatched or future-dated payloads are rejected."),
    ("FR4", "Persist readings and device health.", "A valid reading is stored and last-seen/latency/status fields are updated transactionally."),
    ("FR5", "Handle duplicate delivery safely.", "Repeating a device sequence returns already_processed and does not add a second row."),
    ("FR6", "Display live and historical weather with units.", "Authenticated users see current values in °C, %, hPa, mm, km/h and % battery, plus history."),
    ("FR7", "Manage devices, locations, users and rules.", "Administrators can use Payload Admin; viewers remain read-only."),
    ("FR8", "Detect and manage abnormal conditions.", "Threshold/offline rules create, update, acknowledge, escalate and resolve alerts."),
    ("FR9", "Notify and report.", "An outbox worker supports email preview/provider delivery and reports can be exported."),
    ("FR10", "Demonstrate controlled faults.", "Authorised simulation commands trigger high values, latency, duplicates, invalid data or disconnects with acknowledgements."),
], [1.2, 6.4, 8.0])

doc.add_heading("2.3 Non-Functional Requirements", level=2)
add_table(doc, ["ID", "Quality requirement", "Design response / target"], [
    ("NFR1", "Reliability", "QoS 1, retained status, graceful reconnect and a 10,000-message durable gateway buffer."),
    ("NFR2", "Consistency", "UTC timestamps, positive sequences, UUID message IDs and database uniqueness provide deterministic ordering and idempotency."),
    ("NFR3", "Performance", "Live events should appear without manual refresh; bounded 16 KiB ingestion and indexed queries keep work predictable."),
    ("NFR4", "Security", "Authenticated human sessions, admin/viewer roles, a separate service-account API key, strict schemas and loopback-only development ports."),
    ("NFR5", "Availability", "Independent services and edge buffering allow recovery from temporary application outages; health checks expose infrastructure state."),
    ("NFR6", "Usability", "Responsive navigation, explicit units, status labels, loading/empty/error states and consistent visual hierarchy."),
    ("NFR7", "Maintainability", "Typed TypeScript and Pydantic contracts, migrations, modular services, environment configuration and automated tests."),
    ("NFR8", "Scalability", "Brokered topics and stateless API requests allow more device publishers and web clients; production deployment can replicate consumers."),
], [1.2, 4.2, 10.2])

doc.add_heading("3. System Design", level=1)
doc.add_heading("3.1 High-Level Architecture", level=2)
body(doc, "WeatherGrid is a layered distributed system. Producers, broker, gateway, application, database, workers and clients are separate processes with explicit network and data contracts. This reduces direct dependencies and allows each concern to fail or scale independently. Figure 1 shows the implemented architecture.")
doc.add_picture(str(ROOT / "weathergrid-high-level-architecture-generated.png"), width=Inches(6.7))
p = doc.add_paragraph("Figure 1. WeatherGrid high-level distributed system architecture (created by the project team).", style="Figure Caption")
p.alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_heading("3.2 Components and Communication", level=2)
add_table(doc, ["Layer", "Implemented component", "Responsibility and communication"], [
    ("Device", "Three Python 3.11 simulators", "Generate realistic sensor values, validate with Pydantic, preserve sequences and publish JSON."),
    ("Messaging", "Eclipse Mosquitto 2.1.2", "Routes MQTT telemetry, status, commands and acknowledgements without coupling publishers to subscribers."),
    ("Edge", "Python gateway, Paho MQTT, SQLite", "Subscribes, validates topic/payload identity, timestamps, calls HTTPS/HTTP ingestion, buffers and retries failures."),
    ("Application", "Next.js 15, Payload CMS 3, TypeScript", "Authenticates users/services, validates ingestion, executes transactions, evaluates alerts and serves APIs."),
    ("Data", "PostgreSQL 15", "Stores relational device, reading, alert, notification and audit records with constraints and indexes."),
    ("Workers", "Offline and notification workers", "Poll protected state, apply offline rules and deliver queued notifications with bounded retry."),
    ("Client", "React 19 responsive web UI", "Presents overview, live, history, map, devices, alerts, notifications, reports, system and simulation views."),
], [2.0, 4.5, 9.1])

doc.add_heading("3.3 End-to-End Data Flow", level=2)
body(doc, "The normal telemetry flow is: (1) a station samples sensors; (2) the embedded controller creates a versioned JSON payload with a UUID, persistent sequence and UTC timestamp; (3) the MQTT client publishes to weathergrid/devices/{deviceId}/telemetry at QoS 1; (4) Mosquitto routes the message to the gateway; (5) the gateway checks the topic and Pydantic contract, adds gatewayTimestamp and submits the ingestion API using its service-account key; (6) the API checks body size, schema, time ordering, device identity and authorisation; (7) one PostgreSQL transaction creates the reading, updates the device and evaluates alert rules; and (8) SSE notifies open dashboards, which fetch updated values. If HTTP delivery fails temporarily, the gateway writes the complete message to SQLite and retries with bounded exponential backoff. If a duplicate reaches the API, the device/sequence uniqueness constraint converts it into a successful already-processed response instead of a second reading.")

flow_rows = [
    ("1", "Sensor / simulator", "Sample environment and validate ranges"),
    ("2", "MQTT publish", "Versioned JSON, topic identity, QoS 1"),
    ("3", "Mosquitto", "Route message to subscribed gateway"),
    ("4", "Edge gateway", "Validate, timestamp, deliver or buffer"),
    ("5", "Ingestion API", "Authenticate, validate and transact"),
    ("6", "PostgreSQL + alert engine", "Persist, deduplicate, update status and evaluate rules"),
    ("7", "SSE + dashboard", "Signal change and display values with units"),
]
doc.add_picture(str(FLOWCHART), width=Inches(6.7))
p = doc.add_paragraph("Figure 2. Telemetry processing flow from observation to live dashboard (created by the project team).", style="Figure Caption")
p.alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_heading("3.4 Distributed Computing Principles", level=2)
body(doc, "The broker provides spatial decoupling: publishers do not know which gateway or application will consume their readings. The queue provides temporal decoupling: the central server can be unavailable while the edge continues accepting MQTT messages. At-least-once transport is reconciled with application-level exactly-once persistence through idempotency. PostgreSQL transactions maintain atomic changes across readings and device summaries. SSE implements one-way fan-out to many browser clients with lower overhead than polling for each live update. Independent workers move slow offline checks and notifications out of the request path. These decisions illustrate message-oriented middleware, fault tolerance, replication-ready stateless services, concurrency control and eventual delivery, which are central distributed-computing concepts.")

doc.add_heading("4. Prototype Implementation", level=1)
doc.add_heading("4.1 Technology Stack", level=2)
add_table(doc, ["Technology", "Full form / role", "Use in WeatherGrid"], [
    ("IoT", "Internet of Things", "Connects distributed sensing nodes to software services."),
    ("ESP32 (future hardware)", "Espressif Systems 32-bit microcontroller", "Target controller with Wi-Fi for physical sensors; simulated in the current prototype."),
    ("MQTT", "Message Queuing Telemetry Transport", "Lightweight publish/subscribe messaging through Mosquitto."),
    ("QoS", "Quality of Service", "Level 1 requests at-least-once delivery."),
    ("HTTP/HTTPS", "Hypertext Transfer Protocol / Secure", "Gateway ingestion and browser/API communication."),
    ("API", "Application Programming Interface", "Typed endpoints for telemetry, dashboards, alerts, reports and simulation."),
    ("SSE", "Server-Sent Events", "Pushes change notifications to live browser clients."),
    ("CMS", "Content Management System", "Payload provides collections, authentication, access control and administration."),
    ("SQL", "Structured Query Language", "PostgreSQL persistence and SQLite edge queue."),
    ("UUID", "Universally Unique Identifier", "Globally unique telemetry and command identifiers."),
], [3.5, 5.7, 7.0])

doc.add_heading("4.2 Device and Sensor Layer", level=2)
body(doc, "The software prototype simulates a realistic station rather than claiming that physical sensors are attached. A future ESP32 node would read a digital temperature/humidity sensor such as DHT22 or BME280, a BME280 pressure sensor, a tipping-bucket rain gauge and a cup anemometer; battery voltage would be measured through an analogue-to-digital converter. The controller would poll each interface at a configured interval, convert raw values to engineering units, attach identifiers and publish through Wi-Fi. In the current implementation, independent Python processes produce bounded values for Sydney, Melbourne and Brisbane using device profiles and scenarios. Pydantic rejects invalid types, ranges or timestamps before publication. Sequence state is atomically written outside the source tree so process restarts do not silently reuse numbers.")
code_block(doc, '''{
  "schemaVersion": 1,
  "messageId": "3dd73cd4-2080-40bb-bfa9-e53f7ee493e5",
  "deviceId": "WX-SYD-001",
  "sequenceNumber": 25,
  "temperature": 18.2, "humidity": 64.7,
  "pressure": 1015.3, "rainfall": 0.0,
  "windSpeed": 11.8, "battery": 99.9,
  "deviceTimestamp": "2026-09-24T02:34:38.731Z"
}''', "Code Extract 1. Versioned MQTT telemetry contract (representative payload).")

doc.add_heading("4.3 Messaging and Edge Gateway", level=2)
body(doc, "Telemetry is non-retained because historical storage belongs in PostgreSQL; connectivity status is retained so a new subscriber immediately learns the last online/offline state. Each client also configures an MQTT Last Will for unexpected disconnection. Command and acknowledgement topics enable the application to control simulation scenarios without directly calling a device process. The gateway subscribes with a regular expression-constrained topic, rejects a payload whose device ID disagrees with its topic, and never logs credentials. A temporary API or database failure adds the enriched payload to a bounded SQLite queue. The oldest due item is retried first, with delays capped at 60 seconds. Authentication and permanent validation failures are separated from transient failures so bad data is not retried forever.")
code_block(doc, '''result = client.deliver(telemetry)
if result.status in {ACCEPTED, ALREADY_PROCESSED}:
    return ProcessingResult(result.status)
if result.status == AUTH_FAILURE:
    raise GatewayAuthenticationError()
buffer.enqueue(telemetry)
return ProcessingResult("buffered")''', "Code Extract 2. Simplified gateway delivery and durable-buffer decision.")

doc.add_heading("4.4 Application, Data and User Interface", level=2)
body(doc, "Next.js combines server-rendered React pages and server endpoints in one TypeScript application. Payload CMS defines collections for users, locations, devices, weather readings, alert rules, alerts, simulation commands, service accounts, notification preferences, outbox entries and system events. PostgreSQL migrations version these schemas. Ingestion has a strict 16 KiB body limit, permits no unknown fields, rejects future timestamps and requires an active edge-gateway service account. Human users authenticate separately as admin or viewer. Administrators can configure operational data and resolve alerts; viewers can monitor without mutation privileges.")
body(doc, "The dashboard provides overview, live monitoring, history, devices, map, alerts, notifications, reports, system and simulation pages. Measurements display explicit units: temperature in degrees Celsius, humidity in percent, atmospheric pressure in hectopascals, rainfall in millimetres, wind speed in kilometres per hour and battery charge in percent. Live SSE messages trigger data refresh while a timed fallback refresh protects against a broken stream. Figure 3 shows the running live-monitoring view, not a design mock-up.")
doc.add_picture(str(ROOT / "weathergrid-live-monitoring.png"), width=Inches(6.65))
p = doc.add_paragraph("Figure 3. Authenticated WeatherGrid live-monitoring prototype with three active stations (captured 24 September 2026).", style="Figure Caption")
p.alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_heading("4.5 Alerting, Notifications and Reporting", level=2)
body(doc, "The alert engine evaluates enabled rules after accepted telemetry. Seeded rules cover high temperature, high humidity, abnormal pressure, low battery, high latency and offline devices. Scope can be global, by location or by device. One unresolved alert is permitted per device/rule; repeated breaches update the existing alert, while recovery can resolve it automatically. Administrators can acknowledge or manually resolve an alert with a reason, producing an audit event. A separate worker consumes a notification outbox and either creates a preview or sends through the configured email provider with bounded retries. Report endpoints export filtered operational data and record report activity.")

doc.add_heading("5. Testing and Evaluation", level=1)
doc.add_heading("5.1 Test Strategy and Results", level=2)
body(doc, "Testing follows the architecture rather than relying only on visual inspection. Unit tests cover telemetry validation, alert boundaries, simulation contracts, gateway queue behaviour and processing decisions. Static checks detect TypeScript contract and style defects. Infrastructure health checks verify that PostgreSQL accepts connections and Mosquitto can publish at QoS 1. Live observation verifies that three processes publish concurrently and that their latest values reach persistent storage and the authenticated dashboard.")
add_table(doc, ["Test area", "Method", "Observed result", "Status"], [
    ("Web/business logic", "Vitest suite", "3 files; 8 tests passed in 167 ms", "Pass"),
    ("Edge gateway", "Pytest suite", "7 tests passed in 0.14 s", "Pass"),
    ("Device simulator", "Pytest suite", "27 tests passed in 0.30 s", "Pass"),
    ("Type safety", "TypeScript compiler --noEmit", "No type errors", "Pass"),
    ("Code quality", "ESLint full project", "No lint errors", "Pass"),
    ("Infrastructure", "Docker Compose health checks", "PostgreSQL and Mosquitto healthy", "Pass"),
    ("Concurrent telemetry", "Query latest records", "Sydney, Melbourne and Brisbane reported at the same interval", "Pass"),
    ("Persistence", "Database count on 24 Sep 2026", "4,531 readings, 3 devices, 3 locations and 53 alert records", "Pass"),
    ("Live UI", "Authenticated browser inspection", "Latest station values and units rendered; evidence in Figure 3", "Pass"),
], [3.0, 4.5, 7.0, 1.7])

doc.add_heading("5.2 Critical Behaviour Evaluation", level=2)
body(doc, "Validation tests exercise missing/extra fields, invalid ranges and malformed timestamps. Gateway tests confirm invalid topics, payload/topic device mismatch, accepted delivery, temporary buffering, authentication failure and retry scheduling. Simulator tests cover configuration, bounded weather generation, persisted sequences, MQTT topics and status behaviour. Alert tests check inclusive warning/critical boundaries and recovery. Simulation tests verify supported command types and expiry rules. Together, these tests demonstrate more than a happy path: they target the boundaries where distributed systems commonly fail.")
body(doc, "The duplicate strategy is especially important. MQTT QoS 1 means the broker can redeliver after a lost acknowledgement. WeatherGrid does not assume that transport acknowledgements imply unique application records. Instead, it treats device ID plus sequence number as a stable idempotency key. The API returns success for an already processed sequence, allowing the gateway to remove it from the retry queue. This prevents repeated measurements and alert evaluation while preserving at-least-once reliability.")

doc.add_heading("5.3 Performance and Limitations", level=2)
body(doc, "The local verification stored more than 4,500 readings while all three stations remained active. Unit suites complete in well under one second each, indicating that core validation and decision logic are lightweight. The dominant real-time latency is expected to be network and persistence time rather than sensor computation. The design records device and gateway timestamps, allowing ingestion latency to be measured and used by a high-latency alert. SSE avoids separate full-data polling loops for every browser event, although each change currently causes the client to fetch updated dashboard data. This is appropriate for the three-device prototype but should be load-tested before a large rollout.")
body(doc, "Evaluation is limited by the local academic environment. Simulator accuracy demonstrates software behaviour, not meteorological accuracy; physical sensors require calibration against reference instruments. Local services share one computer, so the tests do not reproduce Internet packet loss, regional failure or high concurrency. Email defaults to preview mode. Security is intentionally local: Mosquitto is loopback-only and development credentials are unsuitable for production. Consequently, the results establish functional correctness and fault-handling logic, but not production capacity or scientific-grade measurement quality.")

doc.add_heading("6. Challenges and Solutions", level=1)
add_table(doc, ["Challenge", "Cause / risk", "Implemented solution and learning"], [
    ("At-least-once duplicates", "QoS 1 may redeliver after acknowledgement loss.", "Persistent sequences, message UUIDs and a database unique constraint make ingestion idempotent. Transport reliability must be paired with application semantics."),
    ("Temporary server outage", "A gateway HTTP request can fail while devices continue publishing.", "SQLite stores enriched payloads durably and retries oldest-first with bounded backoff, preventing short outages from becoming data loss."),
    ("Unsafe or malformed telemetry", "Distributed boundaries cannot trust message format, source identity or time.", "Pydantic and Zod validate strict schemas; the API checks topic/device match, size, ranges, timestamp ordering and authorised service identity."),
    ("False duplicate sequences", "Simulator state originally existed under a source-watched path and restarts could disrupt development behaviour.", "State moved to a runtime directory with atomic writes. This isolates operational state from application source watching."),
    ("Human and machine identity", "Using an administrator credential in the gateway would give excessive privilege.", "A dedicated edge-gateway service account and API key separate machine ingestion from admin/viewer sessions."),
    ("Detecting silent devices", "No telemetry event arrives when power or connectivity is lost.", "A scheduled, advisory-lock-protected worker compares lastSeen against the publish interval, excludes maintenance devices and raises offline alerts."),
    ("Clear interpretation", "Raw numbers without context confused users.", "The live interface labels each measurement with its engineering unit and provides health, freshness and alert status."),
], [4.0, 5.3, 7.1])
body(doc, "These challenges changed the project from a simple sensor dashboard into a distributed system. The main learning was that reliable messaging is not one feature: it requires identity, validation, ordering, durable state, retry policy, idempotency, observability and clear user feedback across multiple components.")

doc.add_heading("7. Conclusion and Future Work", level=1)
body(doc, "WeatherGrid achieved its core objectives. It implements three independent weather nodes, brokered MQTT communication, a validating and buffering edge gateway, transactional PostgreSQL storage, role-based administration, live and historical web views, rule-based alerts, notification processing, reports and controlled failure simulation. The prototype applies distributed computing concepts visibly rather than only describing them: producers and consumers are decoupled, transient failures are buffered, duplicate delivery is made idempotent, background work is separated, and live events are fanned out to authenticated clients. Reproducible test results and live database evidence show that the implemented path works end to end.")
body(doc, "The present system remains a prototype. Its devices are simulated, services run on one development machine, MQTT transport is not TLS-protected, and performance has not been tested at city scale. The highest-priority next step is to build an ESP32 station with calibrated BME280/DHT22, rain-gauge and anemometer inputs and compare measurements with an official reference station. Production deployment should use MQTT over TLS, per-device certificates and access-control lists, secret rotation, encrypted HTTPS, automated backups, monitoring and high-availability PostgreSQL. Kubernetes or managed container services could replicate the stateless application and gateway consumers; broker clustering and partition-aware load tests would quantify scale. Further improvements include offline-first device storage, over-the-air firmware updates, calibration metadata, geospatial interpolation, predictive anomaly detection, push notifications and a progressive web application. These extensions preserve the existing contracts while moving WeatherGrid from a strong academic prototype toward an operational environmental monitoring platform.")

doc.add_heading("References", level=1)
refs = [
    "Banks, A., Briggs, E., Borgendale, K., & Gupta, R. (Eds.). (2019). MQTT version 5.0 (OASIS Standard). OASIS Open. https://docs.oasis-open.org/mqtt/mqtt/v5.0/os/mqtt-v5.0-os.html",
    "Gubbi, J., Buyya, R., Marusic, S., & Palaniswami, M. (2013). Internet of Things (IoT): A vision, architectural elements, and future directions. Future Generation Computer Systems, 29(7), 1645-1660. https://doi.org/10.1016/j.future.2013.01.010",
    "Hernandez, W., Mendez, A., Leiva, A., & Bonilla, J. (2024). Development of a unified IoT platform for assessing meteorological and air quality data in a tropical environment. Sensors, 24(9), 2729. https://doi.org/10.3390/s24092729",
    "Kodali, R. K., & Mandal, S. (2016). IoT based weather station. In 2016 International Conference on Control, Instrumentation, Communication and Computational Technologies (ICCICCT) (pp. 680-683). IEEE. https://doi.org/10.1109/ICCICCT.2016.7988038",
    "Kodali, R. K., & Sahu, A. (2016). An IoT based weather information prototype using WeMos. In 2016 2nd International Conference on Contemporary Computing and Informatics (IC3I). IEEE. https://doi.org/10.1109/IC3I.2016.7918036",
    "Satyanarayanan, M. (2017). The emergence of edge computing. Computer, 50(1), 30-39. https://doi.org/10.1109/MC.2017.9",
]
for ref in refs:
    p = doc.add_paragraph(ref)
    p.paragraph_format.left_indent = Cm(1.27)
    p.paragraph_format.first_line_indent = Cm(-1.27)
    p.paragraph_format.space_after = Pt(8)

doc.add_heading("Appendix A: 15-Minute Presentation and Demonstration Plan", level=1)
add_table(doc, ["Time", "Presenter", "Content / action", "Evidence on screen"], [
    ("0:00-1:30", "Member 1", "Problem, users and project objectives", "Title plus concise problem statement"),
    ("1:30-4:00", "Member 2", "Architecture and distributed principles", "Figure 1; trace device, MQTT, gateway, API, database and SSE"),
    ("4:00-5:30", "Member 3", "Technology and security decisions", "Stack, roles, API key, validation and idempotency"),
    ("5:30-10:30", "All", "Working prototype demonstration", "Login; Overview; Live; trigger a high-temperature simulation; observe alert; acknowledge/resolve; inspect History/Report"),
    ("10:30-12:00", "Member 2", "Failure demonstration", "Briefly stop application delivery or explain queued message; show SQLite retry and recovered reading"),
    ("12:00-13:30", "Member 3", "Testing and measured evidence", "42 passing tests, healthy services, stored readings and limitations"),
    ("13:30-15:00", "Member 1", "Conclusion and future work", "Objectives achieved, ESP32/TLS/calibration roadmap"),
], [2.0, 2.4, 7.3, 5.0])
body(doc, "Demonstration safety: start PostgreSQL and Mosquitto, then the Next.js app, gateway, simulator and workers before class. Confirm three devices are online, open the authenticated dashboard in advance, and keep Figure 3 plus a short screen recording as fallback evidence. Do not display .env, API keys or passwords. Rehearse command expiry and reset-normal so the demonstration returns to a clean state.")

doc.add_heading("Appendix B: Group Contribution Record", level=1)
add_table(doc, ["Member name / ID", "Primary responsibilities", "Evidence (commits, tests, sections)", "Contribution %"], [
    ("[INSERT]", "[INSERT]", "[INSERT]", "[INSERT]"),
    ("[INSERT]", "[INSERT]", "[INSERT]", "[INSERT]"),
    ("[INSERT]", "[INSERT]", "[INSERT]", "[INSERT]"),
], [4.0, 5.0, 6.0, 2.0])

# Keep headings and captions together with following content where possible.
for paragraph in doc.paragraphs:
    if paragraph.style.name.startswith("Heading"):
        paragraph.paragraph_format.keep_with_next = True

doc.core_properties.title = "WeatherGrid: Distributed IoT Weather Monitoring System"
doc.core_properties.subject = "ICT304 Assessment 3 Group Project Report"
doc.core_properties.keywords = "IoT, distributed computing, MQTT, WeatherGrid, edge gateway"
doc.core_properties.comments = "Generated from the implemented WeatherGrid prototype and verified test results."
doc.save(OUTPUT)
print(OUTPUT)
